from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.workbook.workbook import Workbook

from format_and_validate_cpa_data.domain.models import InconsistenciaRegistro, CorrectionLogEntry
from format_and_validate_cpa_data.sanitizer import sanitize_code, is_valid_sigtap_code, try_autocorrect_zero_prefix, sanitize_cnpj


def _find_column_by_keyword(columns: list[str], keywords: list[str], exclude: list[str] = None) -> str | None:
    """Busca o nome de uma coluna baseado em palavras-chave, ignorando case."""
    if exclude is None:
        exclude = []
        
    for col in columns:
        col_lower = str(col).lower()
        if any(ex in col_lower for ex in exclude):
            continue
        if any(kw in col_lower for kw in keywords):
            return col
    return None


def load_workbook_and_sheet(path: Path) -> tuple[Workbook, pd.DataFrame]:
    """Carrega o arquivo excel via openpyxl e os dados via pandas."""
    wb = openpyxl.load_workbook(path)
    df = pd.read_excel(path, sheet_name="PLANOS OPERATIVOS", dtype=str)
    return wb, df


def load_vinculos_lookup(path: Path) -> dict[str, tuple[str | None, str | None]]:
    """
    Lê a aba VINCULOS e retorna um dicionário de lookup indexado pelo número do vínculo.
    
    Returns:
        dict: {numero_vinculo: (estabelecimento, cnes)}
    """
    try:
        df_vinculos = pd.read_excel(path, sheet_name="VINCULOS", dtype=str)
    except ValueError:
        # Se a aba não existir, retorna lookup vazio
        return {}
        
    columns = df_vinculos.columns.tolist()
    col_vinculo = _find_column_by_keyword(columns, ["vinculo", "vínculo"])
    col_estab   = _find_column_by_keyword(columns, ["estabelecimento", "instituicao", "instituição", "contratada"])
    col_cnes    = _find_column_by_keyword(columns, ["cnes"])
    
    lookup = {}
    for _, row in df_vinculos.iterrows():
        key = str(row[col_vinculo]).strip() if col_vinculo and pd.notna(row[col_vinculo]) else None
        if key:
            lookup[key] = (
                str(row[col_estab]).strip() if col_estab and pd.notna(row[col_estab]) else None,
                str(row[col_cnes]).strip()  if col_cnes  and pd.notna(row[col_cnes])  else None,
            )
    return lookup


def extract_inconsistencies(
    df: pd.DataFrame, 
    vinculos_lookup: dict[str, tuple[str | None, str | None]] | None = None,
    existing_codes: set[str] | None = None
) -> list[InconsistenciaRegistro]:
    """Extrai os registros que não possuem um código SIGTAP de 10 dígitos."""
    columns = df.columns.tolist()
    
    # Mapeamento adaptativo
    col_codigo = _find_column_by_keyword(columns, ["codigo", "código", "sigtap"])
    if not col_codigo:
        raise ValueError("Não foi possível encontrar a coluna de código SIGTAP na planilha.")
        
    col_vinculo = _find_column_by_keyword(columns, ["vinculo", "vínculo"])
    col_aditivo = _find_column_by_keyword(columns, ["aditivo"])
    col_plano = _find_column_by_keyword(columns, ["plano operativo"])
    col_descricao = _find_column_by_keyword(columns, ["descricao", "descrição", "procedimento"], exclude=["codigo", "código"])
    col_quantidade = _find_column_by_keyword(columns, ["quantidade", "mensal", "pactuada"])

    inconsistencies = []
    
    # openpyxl: row 1 is header, row 2 is index 0
    # we need the index of the column to update it later, let's find it 
    # but for InconsistenciaRegistro we just need the excel row number
    
    for idx, row in df.iterrows():
        raw_code = row[col_codigo] if pd.notna(row[col_codigo]) else None
        sanitized = sanitize_code(raw_code)
        
        if not is_valid_sigtap_code(sanitized) or (existing_codes is not None and sanitized not in existing_codes):
            vinculo_val = str(row[col_vinculo]).strip() if col_vinculo and pd.notna(row[col_vinculo]) else "N/A"
            
            instituicao, cnes = None, None
            if vinculos_lookup and vinculo_val in vinculos_lookup:
                instituicao, cnes = vinculos_lookup[vinculo_val]
                
            inconsistencies.append(
                InconsistenciaRegistro(
                    linha_excel=int(idx) + 2,  # type: ignore (idx is int)
                    instituicao=instituicao,
                    cnes=cnes,
                    plano_operativo=str(row[col_plano]) if col_plano and pd.notna(row[col_plano]) else "N/A",
                    vinculo=vinculo_val,
                    aditivo=str(row[col_aditivo]) if col_aditivo and pd.notna(row[col_aditivo]) else None,
                    codigo_atual=sanitized,
                    descricao_planilha=str(row[col_descricao]) if col_descricao and pd.notna(row[col_descricao]) else None,
                    quantidade_mensal_pactuada=str(row[col_quantidade]) if col_quantidade and pd.notna(row[col_quantidade]) else None,
                )
            )
            
    return inconsistencies

def pre_sanitize_all_codes(df: pd.DataFrame) -> dict[int, str]:
    """
    Etapa 0: Percorre TODAS as linhas e substitui o valor bruto da célula
    de código pelo resultado de sanitize_code() — somente dígitos.

    Modifica o DataFrame in-place. Retorna um dicionário
    {linha_excel: valor_sanitizado} contendo APENAS as linhas cujo valor
    mudou (ou seja, as células que tinham "lixo" junto ao código).

    Isso garante que o openpyxl escreva valores limpos antes de qualquer
    outra lógica de validação/correção.

    Returns:
        dict[int, str]: {linha_excel: codigo_sanitizado} das células alteradas.

    Raises:
        ValueError: Se a coluna de código não for encontrada.
    """
    columns = df.columns.tolist()
    col_codigo = _find_column_by_keyword(columns, ["codigo", "código", "sigtap"])

    if not col_codigo:
        raise ValueError("Não foi possível encontrar a coluna de código SIGTAP.")

    changed: dict[int, str] = {}

    for idx, row in df.iterrows():
        raw = row[col_codigo]
        raw_str = "" if pd.isna(raw) else str(raw).strip()
        sanitized = sanitize_code(raw)

        # Só registra mudança se o valor realmente alterou
        if sanitized != raw_str:
            df.at[idx, col_codigo] = sanitized
            changed[int(idx) + 2] = sanitized  # linha_excel = pandas_idx + 2

    return changed

def sanitize_instituicoes_cnpj(df: pd.DataFrame) -> dict[int, str]:
    """
    Percorre a aba INSTITUICOES e higieniza a coluna CNPJ.
    Modifica o DataFrame in-place e retorna as correções.
    
    Returns:
        dict[int, str]: {linha_excel: cnpj_sanitizado} das células alteradas.
    """
    columns = df.columns.tolist()
    col_cnpj = _find_column_by_keyword(columns, ["cnpj"])

    if not col_cnpj:
        return {}

    changed: dict[int, str] = {}

    for idx, row in df.iterrows():
        raw = row[col_cnpj]
        raw_str = "" if pd.isna(raw) else str(raw).strip()
        sanitized = sanitize_cnpj(raw)

        # Só registra se houve mudança
        if sanitized != raw_str:
            df.at[idx, col_cnpj] = sanitized
            changed[int(idx) + 2] = sanitized

    return changed

def apply_auto_corrections(
    df: pd.DataFrame,
    existing_codes: set[str] | None = None,
    vinculos_lookup: dict[str, tuple[str | None, str | None]] | None = None
) -> tuple[dict[int, str], int, list[CorrectionLogEntry]]:
    """
    Varre o DataFrame em busca de códigos inválidos e aplica regras de auto-correção.
    Modifica o DataFrame in-place para que as etapas subsequentes leiam os dados corrigidos.
    
    Returns:
        Um dicionário {linha_excel: codigo_corrigido} para ser salvo posteriormente.
        O número total de auto-correções efetuadas.
        Uma lista de CorrectionLogEntry para auditoria.
    """
    columns = df.columns.tolist()
    col_codigo = _find_column_by_keyword(columns, ["codigo", "código", "sigtap"])
    
    if not col_codigo:
        return {}, 0, []
        
    col_vinculo = _find_column_by_keyword(columns, ["vinculo", "vínculo"])
    col_aditivo = _find_column_by_keyword(columns, ["aditivo"])
        
    corrections = {}
    auto_count = 0
    logs = []
    
    for idx, row in df.iterrows():
        raw_code = row[col_codigo] if pd.notna(row[col_codigo]) else None
        sanitized = sanitize_code(raw_code)
        
        if not is_valid_sigtap_code(sanitized) or (existing_codes is not None and sanitized not in existing_codes):
            corrected = try_autocorrect_zero_prefix(sanitized)
            if corrected and (existing_codes is None or corrected in existing_codes):
                df.at[idx, col_codigo] = corrected
                corrections[int(idx) + 2] = corrected  # type: ignore
                auto_count += 1
                
                vinculo_val = str(row[col_vinculo]).strip() if col_vinculo and pd.notna(row[col_vinculo]) else "N/A"
                instituicao, cnes = None, None
                if vinculos_lookup and vinculo_val in vinculos_lookup:
                    instituicao, cnes = vinculos_lookup[vinculo_val]
                    
                aditivo_val = str(row[col_aditivo]) if col_aditivo and pd.notna(row[col_aditivo]) else None
                
                logs.append(
                    CorrectionLogEntry(
                        linha_excel=int(idx) + 2,
                        instituicao=instituicao,
                        vinculo=vinculo_val,
                        cnes=cnes,
                        aditivo=aditivo_val,
                        codigo_original=sanitized,
                        codigo_corrigido=corrected,
                        modo_correcao="Auto-Correção (Omissão de Zero)"
                    )
                )
                
    return corrections, auto_count, logs



def apply_corrections(wb: Workbook, sheet_name: str, col_name: str, corrections: dict[int, str]) -> None:
    """Aplica as correções no objeto openpyxl em memória."""
    sheet = wb[sheet_name]
    
    # Descobre o índice da coluna (1-based no openpyxl)
    header_row = sheet[1]
    col_idx = None
    for cell in header_row:
        if cell.value == col_name:
            col_idx = cell.column
            break
            
    if not col_idx:
        raise ValueError(f"Coluna {col_name} não encontrada no cabeçalho (linha 1).")
        
    for linha_excel, codigo_correto in corrections.items():
        sheet.cell(row=linha_excel, column=col_idx).value = codigo_correto


def save_workbook(wb: Workbook, output_path: Path) -> None:
    """Salva o workbook no disco."""
    wb.save(output_path)


def freeze_formulas_as_values(wb: Workbook, sheet_name: str, df: pd.DataFrame) -> None:
    """Substitui fórmulas por seus valores calculados (max 2 casas decimais) para evitar células em branco."""
    sheet = wb[sheet_name]
    
    for row_idx, row in enumerate(sheet.iter_rows(), start=1):
        if row_idx == 1:
            continue
            
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                try:
                    # O pandas lê os valores cacheados originais, o que nos permite extraí-los.
                    static_value = df.iat[row_idx - 2, cell.column - 1]
                    if isinstance(static_value, float):
                        static_value = round(static_value, 2)
                    cell.value = static_value
                except (IndexError, ValueError):
                    pass

